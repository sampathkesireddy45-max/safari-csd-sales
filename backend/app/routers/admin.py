import csv
import io
from openpyxl import Workbook
from fastapi import APIRouter, Depends, Query, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import Optional
from sqlalchemy.orm import Session, joinedload, selectinload
from sqlalchemy import func
from .. import models, crud
from ..utils.database import get_db
from ..utils.auth import get_current_active_admin_user
from ..utils.timezone import (
    to_ist,
    format_ist_date,
    format_ist_time,
    format_ist_datetime,
    ist_date_to_utc_bounds,
)
from datetime import datetime, timedelta

router = APIRouter(
    prefix="/admin/dashboard",
    tags=["admin-dashboard"],
    responses={404: {"description": "Not found"}},
    dependencies=[Depends(get_current_active_admin_user)],
)

class AdminOpeningStockUpdateRequest(BaseModel):
    product_id: int
    location_id: int
    opening_stock: int
    date_from: Optional[str] = None
    date_to: Optional[str] = None
    reason: Optional[str] = "Admin manual opening stock adjustment"

class AdminReceivingStockUpdateRequest(BaseModel):
    product_id: int
    location_id: int
    receiving_stock: int
    date_from: Optional[str] = None
    date_to: Optional[str] = None
    reason: Optional[str] = "Admin manual receiving stock adjustment"

class AdminClosingStockUpdateRequest(BaseModel):
    product_id: int
    location_id: int
    closing_stock: int
    date_from: Optional[str] = None
    date_to: Optional[str] = None
    reason: Optional[str] = "Admin manual closing stock adjustment"

class AdminRolloverRequest(BaseModel):
    from_year_month: Optional[str] = None
    to_year_month: Optional[str] = None

def _sale_cost(sale):
    return round(sum(
        (item.product.selling_rate if item.product and item.product.selling_rate is not None
         else item.product.rate_per_pc if item.product and item.product.rate_per_pc is not None else 0) * item.quantity
        for item in sale.items
    ), 2)

def _sales_query(db: Session, employee_id: int = None, location_id: str = None, date_from: str = None, date_to: str = None):
    query = db.query(models.Sale).join(models.Employee).join(models.Location).options(
        joinedload(models.Sale.employee),
        joinedload(models.Sale.location),
        selectinload(models.Sale.items).joinedload(models.SaleItem.product),
    )
    if employee_id:
        query = query.filter(models.Sale.employee_id == employee_id)
    if location_id and str(location_id).strip().lower() not in ["", "all", "0", "overall", "network"]:
        try:
            loc_int = int(str(location_id).strip())
            query = query.filter(models.Sale.location_id == loc_int)
        except ValueError:
            pass
    if date_from and str(date_from).strip():
        try:
            start_utc = ist_date_to_utc_bounds(str(date_from).strip(), is_end_date=False)
            query = query.filter(models.Sale.sale_date >= start_utc)
        except Exception:
            try:
                query = query.filter(models.Sale.sale_date >= datetime.fromisoformat(str(date_from).strip()))
            except Exception:
                pass
    if date_to and str(date_to).strip():
        try:
            end_utc = ist_date_to_utc_bounds(str(date_to).strip(), is_end_date=True)
            query = query.filter(models.Sale.sale_date <= end_utc)
        except Exception:
            try:
                query = query.filter(models.Sale.sale_date < datetime.fromisoformat(str(date_to).strip()) + timedelta(days=1))
            except Exception:
                pass
    return query

def _build_index_wise_sales(all_sales, db: Session, employee_id: int = None, location_id: str = None, all_network_sales = None, date_from: str = None):
    """
    Computes aggregated index-wise performance for products according to FINAL INVENTORY MANAGEMENT FORMULAS:
    - Monthly Opening Stock remains fixed throughout the month.
    - Closing Stock = Opening Stock + Total Received In - Total Dispatched Out - Total Pieces Sold
    - Damages are tracked separately and NOT subtracted from displayed Closing Stock.
    - Manual admin closing edits are saved and never overwritten by automatic recalculation.
    """
    product_sales_map = {}
    for sale in all_sales:
        sale_date_str = format_ist_date(sale.sale_date)
        for item in sale.items:
            pid = item.product_id
            if pid not in product_sales_map:
                product_sales_map[pid] = {
                    "pieces_sold": 0,
                    "transactions_count": 0,
                    "sale_dates": set(),
                }
            product_sales_map[pid]["pieces_sold"] += item.quantity
            product_sales_map[pid]["transactions_count"] += 1
            if sale_date_str:
                product_sales_map[pid]["sale_dates"].add(sale_date_str)

    # Overall network sales
    network_sales = all_network_sales if all_network_sales is not None else all_sales
    network_sales_map = {}
    for sale in network_sales:
        for item in sale.items:
            pid = item.product_id
            network_sales_map[pid] = network_sales_map.get(pid, 0) + item.quantity

    # 1. Overall stock from live inventory across all locations
    overall_inv_rows = db.query(
        models.Inventory.product_id,
        func.coalesce(func.sum(models.Inventory.quantity), 0)
    ).group_by(models.Inventory.product_id).all()
    overall_stock_map = {row[0]: int(row[1] or 0) for row in overall_inv_rows}

    # 2. Store specific stock
    is_specific_store = False
    store_stock_map = {}
    if location_id and str(location_id).strip().lower() not in ["", "all", "0", "overall", "network"]:
        try:
            loc_int = int(str(location_id).strip())
            inv_rows = db.query(models.Inventory.product_id, models.Inventory.quantity).filter(
                models.Inventory.location_id == loc_int
            ).all()
            store_stock_map = {row[0]: int(row[1] or 0) for row in inv_rows}
            is_specific_store = True
        except ValueError:
            pass
    elif employee_id:
        emp = db.query(models.Employee).filter(models.Employee.id == employee_id).first()
        if emp and emp.assigned_location_id:
            inv_rows = db.query(models.Inventory.product_id, models.Inventory.quantity).filter(
                models.Inventory.location_id == emp.assigned_location_id
            ).all()
            store_stock_map = {row[0]: int(row[1] or 0) for row in inv_rows}
            is_specific_store = True

    if not is_specific_store:
        store_stock_map = overall_stock_map.copy()

    # Store specific transfers out and in (Only completed 'RECEIVED' transfers move physical inventory)
    store_transfer_out_map = {}
    store_transfer_in_map = {}
    store_damage_map = {}
    target_store_id = None

    if is_specific_store:
        if location_id and str(location_id).strip().lower() not in ["", "all", "0", "overall", "network"]:
            try:
                target_store_id = int(str(location_id).strip())
            except ValueError:
                pass
        elif employee_id:
            emp = db.query(models.Employee).filter(models.Employee.id == employee_id).first()
            if emp:
                target_store_id = emp.assigned_location_id

        if target_store_id:
            out_rows = db.query(
                models.StockTransfer.product_id,
                func.coalesce(func.sum(models.StockTransfer.dispatched_quantity), 0)
            ).filter(
                models.StockTransfer.from_location_id == target_store_id,
                models.StockTransfer.status == "RECEIVED"
            ).group_by(models.StockTransfer.product_id).all()
            store_transfer_out_map = {row[0]: int(row[1] or 0) for row in out_rows}

            in_rows = db.query(
                models.StockTransfer.product_id,
                func.coalesce(func.sum(models.StockTransfer.received_quantity), 0)
            ).filter(
                models.StockTransfer.to_location_id == target_store_id,
                models.StockTransfer.status == "RECEIVED"
            ).group_by(models.StockTransfer.product_id).all()
            store_transfer_in_map = {row[0]: int(row[1] or 0) for row in in_rows}

            damage_rows = db.query(
                models.DamageReportItem.product_id,
                func.coalesce(func.sum(models.DamageReportItem.quantity), 0)
            ).join(
                models.DamageReport, models.DamageReportItem.damage_report_id == models.DamageReport.id
            ).filter(
                models.DamageReport.location_id == target_store_id,
                models.DamageReport.status != "replaced"
            ).group_by(models.DamageReportItem.product_id).all()
            store_damage_map = {row[0]: int(row[1] or 0) for row in damage_rows}

    # Fleet-wide totals for damages and transfer differences
    overall_damage_rows = db.query(
        models.DamageReportItem.product_id,
        func.coalesce(func.sum(models.DamageReportItem.quantity), 0)
    ).join(
        models.DamageReport, models.DamageReportItem.damage_report_id == models.DamageReport.id
    ).filter(
        models.DamageReport.status != "replaced"
    ).group_by(models.DamageReportItem.product_id).all()
    overall_damage_map = {row[0]: int(row[1] or 0) for row in overall_damage_rows}

    fleet_out_rows = db.query(
        models.StockTransfer.product_id,
        func.coalesce(func.sum(models.StockTransfer.dispatched_quantity), 0)
    ).filter(
        models.StockTransfer.status == "RECEIVED"
    ).group_by(models.StockTransfer.product_id).all()
    fleet_transfer_out_map = {row[0]: int(row[1] or 0) for row in fleet_out_rows}

    fleet_in_rows = db.query(
        models.StockTransfer.product_id,
        func.coalesce(func.sum(models.StockTransfer.received_quantity), 0)
    ).filter(
        models.StockTransfer.status == "RECEIVED"
    ).group_by(models.StockTransfer.product_id).all()
    fleet_transfer_in_map = {row[0]: int(row[1] or 0) for row in fleet_in_rows}

    products = db.query(models.Product).filter(models.Product.is_active == True).order_by(models.Product.index_no, models.Product.name).all()

    current_year_month = date_from[:7] if (date_from and len(date_from) >= 7) else to_ist(datetime.utcnow()).strftime("%Y-%m")
    monthly_records = db.query(models.MonthlyStock).filter(models.MonthlyStock.year_month == current_year_month).all()
    store_monthly_map = {(ms.location_id, ms.product_id): ms for ms in monthly_records if ms.location_id is not None}
    overall_monthly_map = {ms.product_id: ms for ms in monthly_records if ms.location_id is None}
    
    index_wise = []
    total_catalog_pieces = 0
    total_catalog_revenue = 0.0
    total_opening_stock = 0
    total_receiving_stock = 0
    total_closing_stock = 0
    total_overall_opening = 0
    total_overall_receiving = 0
    total_overall_closing = 0
    total_store_opening = 0
    total_store_receiving = 0
    total_store_closing = 0
    total_store_damaged = 0
    total_overall_damaged = 0

    for p in products:
        sales_info = product_sales_map.get(p.id, {"pieces_sold": 0, "transactions_count": 0, "sale_dates": set()})
        pieces_sold = sales_info["pieces_sold"]
        sorted_dates = sorted(list(sales_info["sale_dates"]), reverse=True)
        date_sold_str = ", ".join(sorted_dates) if sorted_dates else ""
        
        selling_rate = p.selling_rate if p.selling_rate is not None else (p.rate_per_pc or 0.0)
        case_pack = p.case_pack if p.case_pack and p.case_pack > 0 else 1
        case_rate = round(selling_rate * case_pack, 2)
        revenue = round(pieces_sold * selling_rate, 2)
        cases_sold = round(pieces_sold / case_pack, 2)

        # Overall fleet numbers:
        overall_sold = network_sales_map.get(p.id, 0)
        overall_damaged = overall_damage_map.get(p.id, 0)
        overall_t_out = fleet_transfer_out_map.get(p.id, 0)
        overall_t_in = fleet_transfer_in_map.get(p.id, 0)

        ms_overall = overall_monthly_map.get(p.id)
        is_overall_receiving_edited = False
        if ms_overall and ms_overall.is_receiving_manually_edited and ms_overall.manual_receiving_stock is not None:
            overall_receiving = ms_overall.manual_receiving_stock
            is_overall_receiving_edited = True
        else:
            overall_receiving = overall_t_in

        if ms_overall:
            overall_opening = ms_overall.opening_stock
            if ms_overall.is_closing_manually_edited and ms_overall.manual_closing_stock is not None:
                overall_closing = ms_overall.manual_closing_stock
            else:
                overall_closing = max(0, overall_opening + overall_receiving - overall_t_out - overall_sold)
        else:
            overall_opening = max(0, overall_stock_map.get(p.id, 0) + overall_sold + overall_t_out - overall_receiving)
            overall_closing = max(0, overall_opening + overall_receiving - overall_t_out - overall_sold)

        # Individual store numbers:
        # Formula: Closing Stock = Opening Stock + Total Received In - Total Dispatched Out - Total Pieces Sold
        # Damages must be recorded separately and must NOT be subtracted from displayed Closing Stock.
        store_transfer_out = store_transfer_out_map.get(p.id, 0)
        store_transfer_in = store_transfer_in_map.get(p.id, 0)
        store_sold = pieces_sold
        store_damaged = store_damage_map.get(p.id, 0)

        ms_store = store_monthly_map.get((target_store_id, p.id)) if target_store_id else None
        is_store_receiving_edited = False
        if ms_store and ms_store.is_receiving_manually_edited and ms_store.manual_receiving_stock is not None:
            store_receiving = ms_store.manual_receiving_stock
            is_store_receiving_edited = True
        else:
            store_receiving = store_transfer_in

        is_closing_manually_edited = False
        if ms_store:
            store_opening = ms_store.opening_stock
            if ms_store.is_closing_manually_edited and ms_store.manual_closing_stock is not None:
                store_closing = ms_store.manual_closing_stock
                is_closing_manually_edited = True
            else:
                store_closing = max(0, store_opening + store_receiving - store_transfer_out - store_sold)
        else:
            store_opening = max(0, store_stock_map.get(p.id, 0) + store_sold + store_transfer_out - store_receiving)
            store_closing = max(0, store_opening + store_receiving - store_transfer_out - store_sold)

        # Active view numbers (based on whether store is selected or overall)
        opening_stock = store_opening if is_specific_store else overall_opening
        receiving_stock = store_receiving if is_specific_store else overall_receiving
        closing_stock = store_closing if is_specific_store else overall_closing

        total_catalog_pieces += pieces_sold
        total_catalog_revenue = round(total_catalog_revenue + revenue, 2)
        total_opening_stock += opening_stock
        total_receiving_stock += receiving_stock
        total_closing_stock += closing_stock
        total_overall_opening += overall_opening
        total_overall_receiving += overall_receiving
        total_overall_closing += overall_closing
        total_store_opening += store_opening
        total_store_receiving += store_receiving
        total_store_closing += store_closing
        total_store_damaged += store_damaged
        total_overall_damaged += overall_damaged

        index_wise.append({
            "product_id": p.id,
            "index_no": p.index_no or "N/A",
            "name": p.name,
            "brand": p.brand or "SAFARI",
            "colour": p.colour or "",
            "category": p.category or "",
            "sub_category": p.sub_category or "",
            "case_pack": case_pack,
            "selling_rate": selling_rate,
            "case_rate": case_rate,
            "opening_stock": opening_stock,
            "receiving_stock": receiving_stock,
            "closing_stock": closing_stock,
            "store_opening_stock": store_opening,
            "store_receiving_stock": store_receiving,
            "store_closing_stock": store_closing,
            "overall_opening_stock": overall_opening,
            "overall_receiving_stock": overall_receiving,
            "overall_closing_stock": overall_closing,
            "store_pieces_sold": store_sold,
            "overall_pieces_sold": overall_sold,
            "total_pieces_sold": pieces_sold,
            "total_cases_sold": cases_sold,
            "total_revenue": revenue,
            "transactions_count": sales_info["transactions_count"],
            "date_sold": date_sold_str,
            "has_sales": pieces_sold > 0,
            "store_damaged": store_damaged,
            "overall_damaged": overall_damaged,
            "is_receiving_manually_edited": is_store_receiving_edited if is_specific_store else is_overall_receiving_edited,
            "is_closing_manually_edited": is_closing_manually_edited,
        })

    # Sort: products with sales first (descending pieces, then revenue), then by index_no
    index_wise.sort(key=lambda x: (-x["total_pieces_sold"], -x["total_revenue"], x["index_no"]))

    return index_wise, {
        "total_indexes": len(products),
        "active_indexes_sold": sum(1 for item in index_wise if item["has_sales"]),
        "total_pieces": total_catalog_pieces,
        "total_revenue": total_catalog_revenue,
        "total_opening_stock": total_opening_stock,
        "total_receiving_stock": total_receiving_stock,
        "total_closing_stock": total_closing_stock,
        "overall_opening_stock": total_overall_opening,
        "overall_receiving_stock": total_overall_receiving,
        "overall_closing_stock": total_overall_closing,
        "store_opening_stock": total_store_opening,
        "store_receiving_stock": total_store_receiving,
        "store_closing_stock": total_store_closing,
        "total_store_damaged": total_store_damaged,
        "total_overall_damaged": total_overall_damaged,
    }

@router.get("/sales")
def get_admin_sales_summary(
    employee_id: Optional[int] = Query(None),
    location_id: Optional[str] = Query(None),
    date_from: Optional[str] = Query(None),
    date_to: Optional[str] = Query(None),
    limit: int = Query(100, ge=1, le=1000),
    db: Session = Depends(get_db),
):
    network_query = _sales_query(db, employee_id=employee_id, date_from=date_from, date_to=date_to)
    all_network_sales = network_query.all()

    query = _sales_query(db, employee_id=employee_id, location_id=location_id, date_from=date_from, date_to=date_to)
    sales = query.order_by(models.Sale.sale_date.desc()).limit(limit).all()
    all_sales = query.all()

    index_wise_sales, index_summary = _build_index_wise_sales(
        all_sales, db, employee_id=employee_id, location_id=location_id, all_network_sales=all_network_sales, date_from=date_from
    )

    all_locations = db.query(models.Location).filter(models.Location.is_active == True).order_by(models.Location.name).all()
    is_specific_store = False
    selected_loc_name = "All Stores Combined"
    selected_loc_id = None
    if location_id and str(location_id).strip().lower() not in ["", "all", "0", "overall", "network"]:
        try:
            loc_int = int(str(location_id).strip())
            loc_obj = db.query(models.Location).filter(models.Location.id == loc_int).first()
            if loc_obj:
                selected_loc_name = loc_obj.name
                selected_loc_id = loc_obj.id
                is_specific_store = True
        except ValueError:
            pass

    employee_totals = db.query(
        models.Employee.id,
        models.Employee.employee_id,
        models.Employee.name,
        func.count(models.Sale.id).label("sale_count"),
        func.coalesce(func.sum(models.Sale.total_pieces), 0).label("total_pieces"),
    ).outerjoin(models.Sale, models.Sale.employee_id == models.Employee.id)
    if employee_id:
        employee_totals = employee_totals.filter(models.Employee.id == employee_id)
    if location_id and is_specific_store:
        employee_totals = employee_totals.filter(models.Sale.location_id == selected_loc_id)
    if date_from:
        try:
            start_utc = ist_date_to_utc_bounds(date_from, is_end_date=False)
            employee_totals = employee_totals.filter(models.Sale.sale_date >= start_utc)
        except Exception:
            employee_totals = employee_totals.filter(models.Sale.sale_date >= datetime.fromisoformat(date_from))
    if date_to:
        try:
            end_utc = ist_date_to_utc_bounds(date_to, is_end_date=True)
            employee_totals = employee_totals.filter(models.Sale.sale_date <= end_utc)
        except Exception:
            employee_totals = employee_totals.filter(models.Sale.sale_date < datetime.fromisoformat(date_to) + timedelta(days=1))
    employee_totals = employee_totals.group_by(models.Employee.id).order_by(func.sum(models.Sale.total_pieces).desc()).all()
    
    total_pieces = query.with_entities(func.coalesce(func.sum(models.Sale.total_pieces), 0)).scalar() or 0
    sales_cost = round(sum(_sale_cost(sale) for sale in all_sales), 2)
    employee_costs = {}
    for sale in all_sales:
        employee_costs[sale.employee_id] = round(employee_costs.get(sale.employee_id, 0) + _sale_cost(sale), 2)

    return {
        "overall": {"sale_count": query.count(), "total_pieces": int(total_pieces), "total_cost": sales_cost},
        "selected_location": {
            "id": selected_loc_id,
            "name": selected_loc_name,
            "is_all_stores": not is_specific_store,
        },
        "locations": [
            {"id": loc.id, "location_id": loc.location_id, "name": loc.name}
            for loc in all_locations
        ],
        "employees": [
            {"id": row.id, "employee_id": row.employee_id, "name": row.name, "sale_count": row.sale_count, "total_pieces": int(row.total_pieces or 0), "total_cost": employee_costs.get(row.id, 0)}
            for row in employee_totals if row.sale_count
        ],
        "index_wise_sales": index_wise_sales,
        "index_summary": index_summary,
        "sales": [
            {
                "id": sale.id,
                "sale_date": to_ist(sale.sale_date).isoformat() if sale.sale_date else None,
                "formatted_date": format_ist_date(sale.sale_date),
                "formatted_time": format_ist_time(sale.sale_date),
                "formatted_datetime": format_ist_datetime(sale.sale_date),
                "employee_id": sale.employee_id,
                "employee_name": sale.employee.name if sale.employee else "Unknown",
                "employee_code": sale.employee.employee_id if sale.employee else "",
                "location_id": sale.location_id,
                "location_name": sale.location.name if sale.location else "Unknown",
                "total_pieces": int(sale.total_pieces or 0),
                "total_cost": _sale_cost(sale),
                "item_count": len(sale.items),
                "items": [
                    {
                        "id": it.id,
                        "product_id": it.product_id,
                        "index_no": it.product.index_no if it.product else "N/A",
                        "product_name": it.product.name if it.product else "Product",
                        "brand": it.product.brand if it.product else "",
                        "colour": it.product.colour if it.product else "",
                        "quantity": it.quantity,
                        "case_pack": it.product.case_pack if it.product else 1,
                        "unit_price": it.product.selling_rate if (it.product and it.product.selling_rate is not None) else (it.product.rate_per_pc or 0.0) if it.product else 0.0,
                    }
                    for it in sale.items
                ]
            }
            for sale in sales
        ],
    }

@router.post("/opening-stock")
def update_admin_opening_stock(
    req: AdminOpeningStockUpdateRequest,
    db: Session = Depends(get_db),
    current_admin: models.Employee = Depends(get_current_active_admin_user),
):
    product = db.query(models.Product).filter(models.Product.id == req.product_id).first()
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")

    location = db.query(models.Location).filter(models.Location.id == req.location_id).first()
    if not location:
        raise HTTPException(status_code=404, detail="Store location not found")

    # Volume sold at this store during the period
    sales_q = db.query(func.coalesce(func.sum(models.SaleItem.quantity), 0)).join(models.Sale, models.SaleItem.sale_id == models.Sale.id).filter(
        models.SaleItem.product_id == req.product_id,
        models.Sale.location_id == req.location_id,
    )
    if req.date_from:
        try:
            start_utc = ist_date_to_utc_bounds(req.date_from, is_end_date=False)
            sales_q = sales_q.filter(models.Sale.sale_date >= start_utc)
        except Exception:
            sales_q = sales_q.filter(models.Sale.sale_date >= datetime.fromisoformat(req.date_from))
    if req.date_to:
        try:
            end_utc = ist_date_to_utc_bounds(req.date_to, is_end_date=True)
            sales_q = sales_q.filter(models.Sale.sale_date <= end_utc)
        except Exception:
            sales_q = sales_q.filter(models.Sale.sale_date <= datetime.fromisoformat(req.date_to) + timedelta(days=1))

    pieces_sold = sales_q.scalar() or 0

    damaged_qty = db.query(func.coalesce(func.sum(models.DamageReportItem.quantity), 0)).join(
        models.DamageReport, models.DamageReportItem.damage_report_id == models.DamageReport.id
    ).filter(
        models.DamageReportItem.product_id == req.product_id,
        models.DamageReport.location_id == req.location_id,
        models.DamageReport.status != "replaced"
    ).scalar() or 0

    t_out = db.query(func.coalesce(func.sum(models.StockTransfer.dispatched_quantity), 0)).filter(
        models.StockTransfer.product_id == req.product_id,
        models.StockTransfer.from_location_id == req.location_id,
        models.StockTransfer.status == "RECEIVED"
    ).scalar() or 0

    t_in = db.query(func.coalesce(func.sum(models.StockTransfer.received_quantity), 0)).filter(
        models.StockTransfer.product_id == req.product_id,
        models.StockTransfer.to_location_id == req.location_id,
        models.StockTransfer.status == "RECEIVED"
    ).scalar() or 0

    # Formula: Closing Stock = Opening Stock + Total Received In - Total Dispatched Out - Total Pieces Sold
    # Damages must be recorded separately and must NOT be subtracted from displayed Closing Stock.
    min_required_opening = max(0, pieces_sold + t_out - t_in)
    if req.opening_stock < min_required_opening:
        raise HTTPException(
            status_code=400,
            detail=f"Opening stock ({req.opening_stock} PCS) cannot be less than net volume reduced by sales & transfers ({min_required_opening} PCS) for {product.name} at {location.name}."
        )

    new_closing_stock = req.opening_stock + t_in - t_out - pieces_sold

    year_month = req.date_from[:7] if (req.date_from and len(req.date_from) >= 7) else to_ist(datetime.utcnow()).strftime("%Y-%m")
    ms = db.query(models.MonthlyStock).filter(
        models.MonthlyStock.year_month == year_month,
        models.MonthlyStock.location_id == req.location_id,
        models.MonthlyStock.product_id == req.product_id,
    ).first()

    old_opening = ms.opening_stock if ms else 0
    if not ms:
        ms = models.MonthlyStock(
            year_month=year_month,
            location_id=req.location_id,
            product_id=req.product_id,
            opening_stock=req.opening_stock,
            is_closing_manually_edited=False,
        )
        db.add(ms)
    else:
        ms.opening_stock = req.opening_stock
        if ms.is_closing_manually_edited and ms.manual_closing_stock is not None:
            new_closing_stock = ms.manual_closing_stock

    inv = db.query(models.Inventory).filter(
        models.Inventory.product_id == req.product_id,
        models.Inventory.location_id == req.location_id,
    ).with_for_update().first()

    old_stock = 0
    if not inv:
        inv = models.Inventory(
            product_id=req.product_id,
            location_id=req.location_id,
            quantity=new_closing_stock,
            last_updated=datetime.utcnow()
        )
        db.add(inv)
        db.flush()
    else:
        old_stock = inv.quantity
        inv.quantity = new_closing_stock
        inv.last_updated = datetime.utcnow()

    diff = new_closing_stock - old_stock
    txn = models.InventoryTransaction(
        inventory_id=inv.id,
        quantity_change=diff,
        transaction_type="opening_stock_edit",
        reference_type="admin_opening_stock_adjustment",
        performed_by=current_admin.id
    )
    db.add(txn)

    from ..utils.stock_service import log_admin_stock_audit
    log_admin_stock_audit(
        db=db,
        admin_id=current_admin.id,
        stock_type="STORE_OPENING",
        year_month=year_month,
        location_id=req.location_id,
        product_id=req.product_id,
        previous_value=old_opening,
        updated_value=req.opening_stock,
        reason=req.reason or "Admin manual opening stock adjustment",
    )
    db.commit()

    case_pack = product.case_pack or 1
    return {
        "status": "success",
        "message": f"Successfully updated opening stock for [{product.name}] to {req.opening_stock} PCS at {location.name}. Live closing stock is {new_closing_stock} PCS.",
        "product_id": product.id,
        "product_name": product.name,
        "index_no": product.index_no,
        "opening_stock": req.opening_stock,
        "closing_stock": new_closing_stock,
        "closing_cases": new_closing_stock // case_pack,
        "closing_remainder": new_closing_stock % case_pack,
        "pieces_sold": pieces_sold,
        "location_id": location.id,
        "location_name": location.name
    }

@router.post("/receiving-stock")
def update_admin_receiving_stock(
    req: AdminReceivingStockUpdateRequest,
    db: Session = Depends(get_db),
    current_admin: models.Employee = Depends(get_current_active_admin_user),
):
    """
    Manually edits individual store receiving stock (Received In).
    Any stock received between will be added to receiving stock which will be added to closing stock.
    Admin can edit it and view individual and overall receiving stock.
    """
    if req.receiving_stock < 0:
        raise HTTPException(status_code=400, detail="Receiving stock cannot be negative")

    product = db.query(models.Product).filter(models.Product.id == req.product_id).first()
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")

    location = db.query(models.Location).filter(models.Location.id == req.location_id).first()
    if not location:
        raise HTTPException(status_code=404, detail="Store location not found")

    year_month = req.date_from[:7] if (req.date_from and len(req.date_from) >= 7) else to_ist(datetime.utcnow()).strftime("%Y-%m")

    # Volume sold at this store during the period
    sales_q = db.query(func.coalesce(func.sum(models.SaleItem.quantity), 0)).join(models.Sale, models.SaleItem.sale_id == models.Sale.id).filter(
        models.SaleItem.product_id == req.product_id,
        models.Sale.location_id == req.location_id,
    )
    if req.date_from:
        try:
            start_utc = ist_date_to_utc_bounds(req.date_from, is_end_date=False)
            sales_q = sales_q.filter(models.Sale.sale_date >= start_utc)
        except Exception:
            sales_q = sales_q.filter(models.Sale.sale_date >= datetime.fromisoformat(req.date_from))
    if req.date_to:
        try:
            end_utc = ist_date_to_utc_bounds(req.date_to, is_end_date=True)
            sales_q = sales_q.filter(models.Sale.sale_date <= end_utc)
        except Exception:
            sales_q = sales_q.filter(models.Sale.sale_date <= datetime.fromisoformat(req.date_to) + timedelta(days=1))
    pieces_sold = sales_q.scalar() or 0

    t_out = db.query(func.coalesce(func.sum(models.StockTransfer.dispatched_quantity), 0)).filter(
        models.StockTransfer.product_id == req.product_id,
        models.StockTransfer.from_location_id == req.location_id,
        models.StockTransfer.status == "RECEIVED"
    ).scalar() or 0

    ms = db.query(models.MonthlyStock).filter(
        models.MonthlyStock.year_month == year_month,
        models.MonthlyStock.location_id == req.location_id,
        models.MonthlyStock.product_id == req.product_id,
    ).first()

    old_receiving = ms.manual_receiving_stock if (ms and ms.is_receiving_manually_edited) else 0

    if not ms:
        ms = models.MonthlyStock(
            year_month=year_month,
            location_id=req.location_id,
            product_id=req.product_id,
            opening_stock=0,
            manual_receiving_stock=req.receiving_stock,
            is_receiving_manually_edited=True,
            receiving_edit_reason=req.reason,
            is_closing_manually_edited=False,
        )
        db.add(ms)
    else:
        old_receiving = ms.manual_receiving_stock if ms.is_receiving_manually_edited else 0
        ms.manual_receiving_stock = req.receiving_stock
        ms.is_receiving_manually_edited = True
        ms.receiving_edit_reason = req.reason

    # Formula: Closing Stock = Opening Stock + Receiving Stock - Dispatched Out - Total Pieces Sold
    # If closing stock was manually edited, keep manual closing stock. Otherwise recalculate.
    if ms.is_closing_manually_edited and ms.manual_closing_stock is not None:
        new_closing_stock = ms.manual_closing_stock
    else:
        new_closing_stock = max(0, ms.opening_stock + req.receiving_stock - t_out - pieces_sold)

    inv = db.query(models.Inventory).filter(
        models.Inventory.product_id == req.product_id,
        models.Inventory.location_id == req.location_id,
    ).with_for_update().first()

    old_stock = 0
    if not inv:
        inv = models.Inventory(
            product_id=req.product_id,
            location_id=req.location_id,
            quantity=new_closing_stock,
            last_updated=datetime.utcnow()
        )
        db.add(inv)
        db.flush()
    else:
        old_stock = inv.quantity
        inv.quantity = new_closing_stock
        inv.last_updated = datetime.utcnow()

    diff = new_closing_stock - old_stock
    txn = models.InventoryTransaction(
        inventory_id=inv.id,
        quantity_change=diff,
        transaction_type="receiving_stock_edit",
        reference_type="admin_receiving_stock_adjustment",
        performed_by=current_admin.id
    )
    db.add(txn)

    from ..utils.stock_service import log_admin_stock_audit
    log_admin_stock_audit(
        db=db,
        admin_id=current_admin.id,
        stock_type="STORE_RECEIVING",
        year_month=year_month,
        location_id=req.location_id,
        product_id=req.product_id,
        previous_value=old_receiving,
        updated_value=req.receiving_stock,
        reason=req.reason or "Admin manual receiving stock adjustment",
    )
    db.commit()

    case_pack = product.case_pack or 1
    return {
        "status": "success",
        "message": f"Successfully updated receiving stock for [{product.name}] to {req.receiving_stock} PCS at {location.name}. Live closing stock is {new_closing_stock} PCS.",
        "product_id": product.id,
        "product_name": product.name,
        "index_no": product.index_no,
        "receiving_stock": req.receiving_stock,
        "closing_stock": new_closing_stock,
        "closing_cases": new_closing_stock // case_pack,
        "closing_remainder": new_closing_stock % case_pack,
        "location_id": location.id,
        "location_name": location.name
    }

@router.post("/closing-stock")
def update_admin_closing_stock(
    req: AdminClosingStockUpdateRequest,
    db: Session = Depends(get_db),
    current_admin: models.Employee = Depends(get_current_active_admin_user),
):
    """
    Manually edits individual store closing stock.
    Per Rule 1.B: The admin can manually edit Closing Stock at any time.
    Manual edits must be saved and must not be overwritten automatically by recalculation.
    """
    if req.closing_stock < 0:
        raise HTTPException(status_code=400, detail="Closing stock cannot be negative")

    product = db.query(models.Product).filter(models.Product.id == req.product_id).first()
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")

    location = db.query(models.Location).filter(models.Location.id == req.location_id).first()
    if not location:
        raise HTTPException(status_code=404, detail="Store location not found")

    year_month = req.date_from[:7] if (req.date_from and len(req.date_from) >= 7) else to_ist(datetime.utcnow()).strftime("%Y-%m")

    ms = db.query(models.MonthlyStock).filter(
        models.MonthlyStock.year_month == year_month,
        models.MonthlyStock.location_id == req.location_id,
        models.MonthlyStock.product_id == req.product_id,
    ).first()

    old_closing = 0
    if not ms:
        ms = models.MonthlyStock(
            year_month=year_month,
            location_id=req.location_id,
            product_id=req.product_id,
            opening_stock=0,
            manual_closing_stock=req.closing_stock,
            is_closing_manually_edited=True,
            closing_edit_reason=req.reason,
        )
        db.add(ms)
    else:
        old_closing = ms.manual_closing_stock if ms.is_closing_manually_edited else ms.opening_stock
        ms.manual_closing_stock = req.closing_stock
        ms.is_closing_manually_edited = True
        ms.closing_edit_reason = req.reason

    inv = db.query(models.Inventory).filter(
        models.Inventory.product_id == req.product_id,
        models.Inventory.location_id == req.location_id,
    ).with_for_update().first()

    old_stock = 0
    if not inv:
        inv = models.Inventory(
            product_id=req.product_id,
            location_id=req.location_id,
            quantity=req.closing_stock,
            last_updated=datetime.utcnow()
        )
        db.add(inv)
        db.flush()
    else:
        old_stock = inv.quantity
        inv.quantity = req.closing_stock
        inv.last_updated = datetime.utcnow()

    diff = req.closing_stock - old_stock
    txn = models.InventoryTransaction(
        inventory_id=inv.id,
        quantity_change=diff,
        transaction_type="closing_stock_edit",
        reference_type="admin_closing_stock_adjustment",
        performed_by=current_admin.id
    )
    db.add(txn)

    from ..utils.stock_service import log_admin_stock_audit
    log_admin_stock_audit(
        db=db,
        admin_id=current_admin.id,
        stock_type="STORE_CLOSING",
        year_month=year_month,
        location_id=req.location_id,
        product_id=req.product_id,
        previous_value=old_closing,
        updated_value=req.closing_stock,
        reason=req.reason or "Admin manual closing stock adjustment",
    )
    db.commit()

    case_pack = product.case_pack or 1
    return {
        "status": "success",
        "message": f"Successfully updated manual closing stock for [{product.name}] to {req.closing_stock} PCS at {location.name}.",
        "product_id": product.id,
        "product_name": product.name,
        "index_no": product.index_no,
        "closing_stock": req.closing_stock,
        "closing_cases": req.closing_stock // case_pack,
        "closing_remainder": req.closing_stock % case_pack,
        "location_id": location.id,
        "location_name": location.name
    }

@router.post("/rollover")
def trigger_stock_rollover(
    req: AdminRolloverRequest,
    db: Session = Depends(get_db),
    current_admin: models.Employee = Depends(get_current_active_admin_user),
):
    """
    Month-End Stock Rollover:
    On the first day of each new month (or triggered by admin):
    New Month Opening Stock = Previous Month's Final Closing Stock
    Preserves all transaction records and historical snapshots without overwriting.
    """
    from ..utils.stock_service import get_current_year_month, get_previous_year_month, get_or_create_monthly_stock
    current_ym = get_current_year_month()
    from_ym = req.from_year_month or get_previous_year_month(current_ym)
    to_ym = req.to_year_month or current_ym

    locations = db.query(models.Location).filter(models.Location.is_active == True).all()
    products = db.query(models.Product).filter(models.Product.is_active == True).all()

    rollover_count = 0
    for loc in locations:
        for p in products:
            get_or_create_monthly_stock(db, year_month=to_ym, location_id=loc.id, product_id=p.id)
            rollover_count += 1

    return {
        "status": "success",
        "message": f"Successfully rolled over monthly stock from {from_ym} to {to_ym} for {rollover_count} product-store combinations.",
        "from_year_month": from_ym,
        "to_year_month": to_ym,
        "combinations_count": rollover_count
    }

@router.get("/sales/export")
def export_admin_sales(
    employee_id: int = Query(None),
    location_id: str = Query(None),
    date_from: str = Query(None),
    date_to: str = Query(None),
    db: Session = Depends(get_db),
):
    sales = _sales_query(db, employee_id, location_id, date_from, date_to).order_by(models.Sale.sale_date.desc()).all()
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow([
        "Sale ID",
        "Sale Date (IST)",
        "Sale Time (IST)",
        "Employee ID",
        "Employee",
        "Location",
        "Index & Items Breakdown",
        "Total Pieces",
        "Total Cost (INR)",
        "Line Items"
    ])
    for sale in sales:
        sale_date = format_ist_date(sale.sale_date)
        sale_time = format_ist_time(sale.sale_date)
        items_breakdown = "; ".join(
            f"{it.product.index_no or ''} {it.product.name if it.product else 'Product'} ({it.quantity} pcs @ ₹{it.product.selling_rate if it.product and it.product.selling_rate is not None else 0})"
            for it in sale.items
        )
        writer.writerow([
            sale.id,
            sale_date,
            sale_time,
            sale.employee.employee_id if sale.employee else "",
            sale.employee.name if sale.employee else "Unknown",
            sale.location.name if sale.location else "Unknown",
            items_breakdown,
            sale.total_pieces or 0,
            f"{_sale_cost(sale):.2f}",
            len(sale.items),
        ])
    output.seek(0)
    csv_content = "\ufeff" + output.getvalue()
    scope = f"location_{location_id}" if location_id and location_id not in ["all", "0"] else (f"employee_{employee_id}" if employee_id else "overall")
    return StreamingResponse(
        iter([csv_content]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=sales_transactions_{scope}.csv"},
    )

@router.get("/sales/export/index-wise/csv")
def export_admin_sales_index_wise_csv(
    employee_id: int = Query(None),
    location_id: str = Query(None),
    date_from: str = Query(None),
    date_to: str = Query(None),
    db: Session = Depends(get_db),
):
    network_query = _sales_query(db, employee_id, None, date_from, date_to)
    all_network_sales = network_query.all()

    sales = _sales_query(db, employee_id, location_id, date_from, date_to).order_by(models.Sale.sale_date.desc()).all()
    index_wise_sales, _ = _build_index_wise_sales(sales, db, employee_id=employee_id, location_id=location_id, all_network_sales=all_network_sales)
    
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow([
        "Index No.",
        "Product / Model Name",
        "Brand",
        "Colour",
        "Category",
        "Sub Category",
        "Case Pack (Pcs)",
        "Store Opening Stock (Pcs)",
        "Store Receiving Stock (Pcs)",
        "Store Volume Sold (Pcs)",
        "Store Closing Stock (Pcs)",
        "Overall Opening Stock (Pcs)",
        "Overall Receiving Stock (Pcs)",
        "Overall Closing Stock (Pcs)",
        "Selling Rate / Pc (INR)",
        "Case Rate (INR)",
        "Total Cases Sold",
        "Total Revenue (INR)",
        "Date Sold",
    ])
    for it in index_wise_sales:
        writer.writerow([
            it["index_no"],
            it["name"],
            it["brand"],
            it["colour"],
            it["category"],
            it["sub_category"],
            it["case_pack"],
            it.get("store_opening_stock", 0),
            it.get("store_receiving_stock", 0),
            it["total_pieces_sold"],
            it.get("store_closing_stock", 0),
            it.get("overall_opening_stock", 0),
            it.get("overall_receiving_stock", 0),
            it.get("overall_closing_stock", 0),
            f"{it['selling_rate']:.2f}",
            f"{it['case_rate']:.2f}",
            f"{it['total_cases_sold']:.2f}",
            f"{it['total_revenue']:.2f}",
            it.get("date_sold", ""),
        ])
    output.seek(0)
    csv_content = "\ufeff" + output.getvalue()
    scope = f"location_{location_id}" if location_id and location_id not in ["all", "0"] else (f"employee_{employee_id}" if employee_id else "overall")
    return StreamingResponse(
        iter([csv_content]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=index_wise_sales_{scope}.csv"},
    )

@router.get("/sales/export/xlsx")
def export_admin_sales_xlsx(
    employee_id: int = Query(None),
    location_id: str = Query(None),
    date_from: str = Query(None),
    date_to: str = Query(None),
    db: Session = Depends(get_db),
):
    network_query = _sales_query(db, employee_id, None, date_from, date_to)
    all_network_sales = network_query.all()

    sales = _sales_query(db, employee_id, location_id, date_from, date_to).order_by(models.Sale.sale_date.desc()).all()
    index_wise_sales, _ = _build_index_wise_sales(sales, db, employee_id=employee_id, location_id=location_id, all_network_sales=all_network_sales)

    workbook = Workbook()

    # ---------------- SHEET 1: Index-Wise Sales Summary ----------------
    ws_index = workbook.active
    ws_index.title = "Sales Dashboard"
    ws_index.append([
        "Index No.",
        "Product / Model Name",
        "Brand",
        "Colour",
        "Category",
        "Sub Category",
        "Case Pack (Pcs)",
        "Store Opening Stock (Pcs)",
        "Store Receiving Stock (Pcs)",
        "Store Volume Sold (Pcs)",
        "Store Closing Stock (Pcs)",
        "Overall Opening Stock (Pcs)",
        "Overall Receiving Stock (Pcs)",
        "Overall Closing Stock (Pcs)",
        "Selling Rate / Pc (INR)",
        "Case Rate (INR)",
        "Total Cases Sold",
        "Total Revenue (INR)",
        "Date Sold",
    ])

    for it in index_wise_sales:
        ws_index.append([
            it["index_no"],
            it["name"],
            it["brand"],
            it["colour"],
            it["category"],
            it["sub_category"],
            it["case_pack"],
            it.get("store_opening_stock", 0),
            it.get("store_receiving_stock", 0),
            it["total_pieces_sold"],
            it.get("store_closing_stock", 0),
            it.get("overall_opening_stock", 0),
            it.get("overall_receiving_stock", 0),
            it.get("overall_closing_stock", 0),
            it["selling_rate"],
            it["case_rate"],
            it["total_cases_sold"],
            it["total_revenue"],
            it.get("date_sold", ""),
        ])

    total_idx_row = len(index_wise_sales) + 2
    ws_index.append([
        "TOTAL",
        "ALL INDEXES SUMMARY",
        "", "", "", "", "",
        f"=SUM(H2:H{total_idx_row-1})",
        f"=SUM(I2:I{total_idx_row-1})",
        f"=SUM(J2:J{total_idx_row-1})",
        f"=SUM(K2:K{total_idx_row-1})",
        f"=SUM(L2:L{total_idx_row-1})",
        f"=SUM(M2:M{total_idx_row-1})",
        f"=SUM(N2:N{total_idx_row-1})",
        "", "",
        f"=SUM(Q2:Q{total_idx_row-1})",
        f"=SUM(R2:R{total_idx_row-1})",
        "",
    ])

    for cell in ws_index[1]:
        cell.font = cell.font.copy(bold=True)
    for cell in ws_index[total_idx_row]:
        cell.font = cell.font.copy(bold=True)

    ws_index.freeze_panes = "A2"
    ws_index.auto_filter.ref = f"A1:S{total_idx_row-1}"

    # Number formats
    for row in ws_index.iter_rows(min_row=2, max_row=total_idx_row, min_col=8, max_col=14):
        row[0].number_format = '#,##0'
        row[1].number_format = '#,##0'
        row[2].number_format = '#,##0'
        row[3].number_format = '#,##0'
        row[4].number_format = '#,##0'
        row[5].number_format = '#,##0'
        row[6].number_format = '#,##0'
    for row in ws_index.iter_rows(min_row=2, max_row=total_idx_row, min_col=15, max_col=16):
        row[0].number_format = '#,##0.00'
        row[1].number_format = '#,##0.00'
    for row in ws_index.iter_rows(min_row=2, max_row=total_idx_row, min_col=17, max_col=18):
        row[0].number_format = '#,##0.00'
        row[1].number_format = '#,##0.00'

    column_widths = {
        "A": 14, "B": 32, "C": 14, "D": 14, "E": 12, "F": 16,
        "G": 16, "H": 20, "I": 20, "J": 20, "K": 20, "L": 20,
        "M": 22, "N": 18, "O": 18, "P": 22, "Q": 26
    }
    for col, width in column_widths.items():
        ws_index.column_dimensions[col].width = width

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    scope = f"location_{location_id}" if location_id and location_id not in ["all", "0"] else (f"employee_{employee_id}" if employee_id else "overall")
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=sales_{scope}.xlsx"},
    )

@router.get("/stats")
def get_admin_dashboard_stats(db: Session = Depends(get_db)):
    # Total active employees
    total_employees = db.query(models.Employee).filter(models.Employee.status == "active").count()
    # Total active locations
    total_locations = db.query(models.Location).filter(models.Location.is_active == True).count()
    # Total active products
    total_products = db.query(models.Product).filter(models.Product.is_active == True).count()
    # Total stock across all locations
    total_stock_result = db.query(models.Inventory).with_entities(
        func.sum(models.Inventory.quantity)
    ).scalar()
    total_stock = total_stock_result if total_stock_result is not None else 0
    total_in_transit = db.query(func.coalesce(func.sum(models.StockTransfer.quantity), 0)).filter(
        models.StockTransfer.status == "SENT"
    ).scalar() or 0
    total_stock = total_stock + total_in_transit

    ist_today = to_ist(datetime.utcnow()).date()
    start_today = ist_date_to_utc_bounds(str(ist_today), is_end_date=False)
    end_today = ist_date_to_utc_bounds(str(ist_today), is_end_date=True)

    pieces_sold_today = db.query(models.Sale).filter(
        models.Sale.sale_date >= start_today,
        models.Sale.sale_date <= end_today
    ).with_entities(
        func.sum(models.Sale.total_pieces)
    ).scalar() or 0

    # Current month bounds in IST
    start_month_ist = datetime(ist_today.year, ist_today.month, 1)
    if ist_today.month == 12:
        end_month_ist = datetime(ist_today.year + 1, 1, 1) - timedelta(seconds=1)
    else:
        end_month_ist = datetime(ist_today.year, ist_today.month + 1, 1) - timedelta(seconds=1)
    start_month = ist_date_to_utc_bounds(start_month_ist.strftime("%Y-%m-%d"), is_end_date=False)
    end_month = ist_date_to_utc_bounds(end_month_ist.strftime("%Y-%m-%d"), is_end_date=True)

    pieces_sold_this_month = db.query(models.Sale).filter(
        models.Sale.sale_date >= start_month,
        models.Sale.sale_date <= end_month
    ).with_entities(
        func.sum(models.Sale.total_pieces)
    ).scalar() or 0

    # Current year bounds in IST
    start_year = ist_date_to_utc_bounds(f"{ist_today.year}-01-01", is_end_date=False)
    end_year = ist_date_to_utc_bounds(f"{ist_today.year}-12-31", is_end_date=True)

    pieces_sold_this_year = db.query(models.Sale).filter(
        models.Sale.sale_date >= start_year,
        models.Sale.sale_date <= end_year
    ).with_entities(
        func.sum(models.Sale.total_pieces)
    ).scalar() or 0

    total_damaged = db.query(models.DamageReportItem).join(
        models.DamageReport, models.DamageReportItem.damage_report_id == models.DamageReport.id
    ).filter(
        models.DamageReport.status != "replaced"
    ).with_entities(
        func.sum(models.DamageReportItem.quantity)
    ).scalar() or 0

    # Recent 5 damage reports (active, not replaced)
    recent_damages_raw = db.query(models.DamageReport).filter(
        models.DamageReport.status != "replaced"
    ).order_by(models.DamageReport.report_date.desc()).limit(5).all()
    recent_damages = [crud.damage_report.enrich_damage_report(r) for r in recent_damages_raw]

    formatted_damages = []
    for r in recent_damages:
        formatted_damages.append({
            "id": r.id,
            "employee_name": r.employee_name,
            "location_name": r.location_name,
            "report_date": format_ist_datetime(r.report_date),
            "total_quantity": sum(i.quantity for i in r.items),
            "reasons": "; ".join(filter(None, [i.description for i in r.items])),
            "photos_count": len(r.photos),
            "photos": [
                {
                    "id": p.id,
                    "photo_url": p.photo_url,
                    "storage_key": p.storage_key,
                    "mime_type": p.mime_type,
                    "file_size": p.file_size
                }
                for p in r.photos
            ]
        })

    # Real Database 7-Day Trend Analytics - Single fast batch query
    seven_days_ago = ist_today - timedelta(days=6)
    range_start = ist_date_to_utc_bounds(str(seven_days_ago), is_end_date=False)
    range_end = ist_date_to_utc_bounds(str(ist_today), is_end_date=True)

    # 1 query for all sales in the 7-day window
    recent_sales = db.query(models.Sale.sale_date, models.Sale.total_pieces).filter(
        models.Sale.sale_date >= range_start,
        models.Sale.sale_date <= range_end
    ).all()
    sales_by_date = {}
    for s_date, pieces in recent_sales:
        if s_date:
            d_str = to_ist(s_date).strftime("%Y-%m-%d")
            sales_by_date[d_str] = sales_by_date.get(d_str, 0) + (pieces or 0)

    # 1 query for all damages in the 7-day window
    recent_damage_items = db.query(models.DamageReport.report_date, models.DamageReportItem.quantity).join(
        models.DamageReport, models.DamageReportItem.damage_report_id == models.DamageReport.id
    ).filter(
        models.DamageReport.report_date >= range_start,
        models.DamageReport.report_date <= range_end,
        models.DamageReport.status != "replaced"
    ).all()
    damages_by_date = {}
    for r_date, qty in recent_damage_items:
        if r_date:
            d_str = to_ist(r_date).strftime("%Y-%m-%d")
            damages_by_date[d_str] = damages_by_date.get(d_str, 0) + (qty or 0)

    daily_trends = []
    for i in range(6, -1, -1):
        target_date = ist_today - timedelta(days=i)
        d_str = target_date.strftime("%Y-%m-%d")
        daily_trends.append({
            "date": target_date.strftime("%a %d"),
            "fullDate": d_str,
            "sales": int(sales_by_date.get(d_str, 0)),
            "damages": int(damages_by_date.get(d_str, 0)),
        })

    return {
        "totalEmployees": total_employees,
        "totalLocations": total_locations,
        "totalProducts": total_products,
        "totalStock": total_stock,
        "piecesSoldToday": int(pieces_sold_today),
        "piecesSoldThisMonth": int(pieces_sold_this_month),
        "piecesSoldThisYear": int(pieces_sold_this_year),
        "totalDamaged": int(total_damaged),
        "recentDamages": formatted_damages,
        "dailyTrends": daily_trends,
    }